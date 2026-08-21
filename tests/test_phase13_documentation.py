import csv
import unittest
from pathlib import Path

from docx import Document


class FinalDocumentationTest(unittest.TestCase):
    def test_all_phase_validation_records_exist(self):
        self.assertTrue(Path("docs/phase_00_baseline_validation.txt").exists())
        for phase in range(1, 14):
            path = Path(f"docs/phase_{phase:02d}_validation.txt")
            self.assertTrue(path.exists(), path)
            self.assertGreater(path.stat().st_size, 100, path)

    def test_diagrams_and_evaluation_artifacts_are_nonempty(self):
        required = ("system_flow.png", "dfd_level_0.png", "dfd_level_1.png", "dfd_level_2.png",
                    "er_diagram.png", "confusion_matrix.png", "roc_curve.png")
        for filename in required:
            path = Path("docs/artifacts") / filename
            self.assertTrue(path.exists(), path)
            self.assertGreater(path.stat().st_size, 10_000, path)
        with Path("docs/artifacts/model_metrics.csv").open(encoding="utf-8") as handle:
            rows = list(csv.DictReader(handle))
        self.assertEqual({row["model"] for row in rows},
                         {"rules", "isolation_forest", "one_class_svm", "lof", "combined"})

    def test_word_report_opens_and_contains_required_sections_and_figures(self):
        path = Path("docs/Intelligent_Log_Forensics_Final_Report.docx")
        self.assertGreater(path.stat().st_size, 100_000)
        document = Document(path)
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        for phrase in ("Abstract", "System Architecture", "Database Design", "Evaluation and Results",
                       "Security, Trust, and Limitations", "Conclusion and Future Work",
                       "must not be presented as production accuracy"):
            self.assertIn(phrase, text)
        self.assertGreaterEqual(len(document.inline_shapes), 7)
        self.assertGreaterEqual(len(document.tables), 8)


if __name__ == "__main__":
    unittest.main()
