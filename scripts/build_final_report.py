import csv
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
ARTIFACTS = ROOT / "docs/artifacts"
OUTPUT = ROOT / "docs/Intelligent_Log_Forensics_Final_Report.docx"


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def add_table(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = str(header)
        set_cell_shading(cell, "DCEEFF")
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    document.add_paragraph()
    return table


def add_bullets(document, items):
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def add_numbered(document, items):
    for item in items:
        document.add_paragraph(item, style="List Number")


def add_figure(document, filename, caption, width=6.4):
    document.add_picture(str(ARTIFACTS / filename), width=Inches(width))
    paragraph = document.paragraphs[-1]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph = document.add_paragraph(caption)
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.runs[0].italic = True


def heading(document, text, level=1):
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.space_before = Pt(10)
    return paragraph


def build():
    with (ARTIFACTS / "model_metrics.csv").open(encoding="utf-8") as handle:
        metric_rows = list(csv.DictReader(handle))
    combined_metrics = next(row for row in metric_rows if row["model"] == "combined")
    evaluation_summary = dict(
        line.split("=", 1)
        for line in (ARTIFACTS / "evaluation_summary.txt").read_text(encoding="utf-8").splitlines()
        if "=" in line
    )

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(.7)
    section.bottom_margin = Inches(.7)
    section.left_margin = Inches(.8)
    section.right_margin = Inches(.8)
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.color.rgb = RGBColor(16, 32, 51)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("INTELLIGENT LOG FORENSICS").bold = True
    title.runs[0].font.size = Pt(28)
    subtitle = document.add_paragraph("Explainable Multi-Algorithm Log Anomaly Detection,\nLive Incident Correlation, and Forensic Integrity")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(16)
    document.add_paragraph("\nFINAL PROJECT REPORT", style="Title").alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_table(document, ["Submission field", "Value to complete before submission"], [
        ("Student name", "____________________________"), ("Register / roll number", "____________________________"),
        ("Programme / department", "____________________________"), ("Institution", "____________________________"),
        ("Project guide", "____________________________"), ("Academic year", "2025–2026"),
    ], [2.2, 4.2])
    p = document.add_paragraph(f"Repository-generated report date: {date.today().isoformat()}")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_page_break()

    heading(document, "Declaration")
    document.add_paragraph(
        "I declare that the project entitled “Intelligent Log Forensics” is original work carried out "
        "for academic evaluation. External libraries and standards are acknowledged in the references. "
        "The student must review and sign this declaration before submission."
    )
    document.add_paragraph("Student signature: ____________________    Date: ____________________")
    heading(document, "Certificate", 1)
    document.add_paragraph(
        "This is to certify that the project report entitled “Intelligent Log Forensics” has been completed "
        "under the supervision of the undersigned guide and satisfies the applicable institutional review requirements."
    )
    document.add_paragraph("Guide signature: ____________________    Head of Department: ____________________")
    heading(document, "Acknowledgement", 1)
    document.add_paragraph(
        "The author acknowledges the project guide, department faculty, reviewers, and the maintainers of "
        "the open-source tools used in this work. No paid API or subscription is required by the system."
    )
    document.add_page_break()

    heading(document, "Abstract")
    document.add_paragraph(
        "Intelligent Log Forensics is a self-hosted Flask and PostgreSQL application that converts uploaded "
        "or live-generated application logs into normalized evidence, explainable detections, MITRE ATT&CK mappings, "
        "correlated incident timelines, analyst feedback, and downloadable reports. The research component compares "
        "transparent rules with PyOD Isolation Forest, One-Class SVM, Local Outlier Factor, and a PyOD average ensemble. "
        "Training and evaluation use deterministic labeled data emitted by the same synthetic generator schema used at "
        "runtime, avoiding the feature mismatch of unrelated network-flow datasets. Integrity is supported through SHA-256 "
        "hashing, source provenance, parameterized raw SQL, migration tracking, and a documented forensic trust score. "
        f"A held-out {evaluation_summary['held_out_rows']}-row evaluation produced a combined F1 of "
        f"{float(combined_metrics['f1']):.3f} and ROC AUC of {float(evaluation_summary['roc_auc']):.3f} "
        "on the controlled synthetic set."
    )
    heading(document, "Keywords", 2)
    document.add_paragraph("Log forensics; anomaly detection; PyOD; explainable AI; MITRE ATT&CK; PostgreSQL; SSE; SHA-256.")

    heading(document, "Contents")
    add_numbered(document, [
        "Introduction and objectives", "Requirements and scope", "System architecture and data-flow diagrams",
        "Database design", "Detection and machine-learning methodology", "Implementation",
        "Evaluation and results", "Security, trust, and limitations", "Testing and deployment",
        "Conclusion and future work", "References and appendices",
    ])
    document.add_page_break()

    heading(document, "1. Introduction and Objectives")
    document.add_paragraph(
        "Operational logs are heterogeneous, high-volume evidence. Manual inspection is slow, while opaque anomaly "
        "scores are difficult to defend in an examination or incident review. This project provides a reproducible "
        "pipeline in which every detection includes a reason and every database value shown in the UI crosses an observable API boundary."
    )
    heading(document, "1.1 Objectives", 2)
    add_bullets(document, [
        "Normalize CSV, JSON, TXT, and LOG evidence into one application-log schema.",
        "Compare transparent rules with three justified anomaly algorithms and a combined ensemble.",
        "Tag detections with a controlled 16-entry taxonomy and MITRE ATT&CK references.",
        "Correlate evidence into incident timelines and publish committed events live through SSE.",
        "Preserve evidence integrity with SHA-256, provenance, raw SQL migrations, and parameterized queries.",
        "Capture analyst feedback, version retraining artifacts, and measure behavior without promising monotonic improvement.",
        "Run at zero software-license cost using self-hostable open-source dependencies.",
    ])

    heading(document, "2. Requirements and Scope")
    add_table(document, ["Area", "Final requirement"], [
        ("Cost", "Free/open-source/self-hosted; no paid service"),
        ("Persistence", "PostgreSQL via psycopg2 repositories and parameterized raw SQL; no ORM"),
        ("Secrets", "Environment only; excluded from Git and Docker build context"),
        ("Evidence", "In-memory processing with one temporary audit folder and cleanup policy"),
        ("Authentication", "Flask-Login retained; optional cookie JWT API with refresh and rate limits"),
        ("Live UI", "One-way SSE plus browser EventSource; no WebSocket requirement"),
        ("Deployment", "Docker Compose with PostgreSQL and named volumes; Kubernetes parked"),
    ])

    heading(document, "3. System Architecture")
    add_figure(document, "system_flow.png", "Figure 1. Final system flow")
    document.add_paragraph(
        "Uploads and live generator batches converge on process_records. The feature engine provides the identical "
        "numeric schema used by evaluation and ML inference. Repositories own all SQL. Events are published only after "
        "the database transaction commits, preventing the UI from displaying rolled-back evidence."
    )
    heading(document, "3.1 Data-Flow Diagrams", 2)
    add_figure(document, "dfd_level_0.png", "Figure 2. DFD Level 0 — system context", 6.0)
    add_figure(document, "dfd_level_1.png", "Figure 3. DFD Level 1 — major processes")
    add_figure(document, "dfd_level_2.png", "Figure 4. DFD Level 2 — evidence processing")
    heading(document, "3.2 Entity–Relationship Design", 2)
    add_figure(document, "er_diagram.png", "Figure 5. Final ER overview")

    heading(document, "4. Database Design")
    document.add_paragraph(
        "The final schema is defined by database/schema.sql and four ordered migrations. schema_migrations records "
        "each filename in the same transaction as the migration. Existing data is preserved through additive changes "
        "and explicit taxonomy backfill. The detailed table/constraint specification is maintained in docs/database_design.md."
    )
    add_table(document, ["Data group", "Tables"], [
        ("Identity and evidence", "users, uploaded_files, normalized_logs"),
        ("Quality and detection", "data_quality_results, risk_events, mitre_mappings"),
        ("Correlation and reporting", "incidents, incident_events, reports"),
        ("Model governance", "model_retraining_runs, schema_migrations"),
    ])

    heading(document, "5. Detection and ML Methodology")
    heading(document, "5.1 Dataset strategy", 2)
    document.add_paragraph(
        "The LogGenerator emits normal, scan, bruteforce, and breach rows. Ground truth is the generator mode. "
        "The feature vector is failed_login_count, status_code, route_frequency, response_time, error_frequency, "
        "request_rate, user_id, and ip_frequency. This matches runtime normalized application logs and deliberately "
        "does not use CICIDS2017 network-flow fields."
    )
    heading(document, "5.2 Training/evaluation split and metrics", 2)
    document.add_paragraph(
        "A stratified 70/30 split keeps mode proportions stable. Only normal training rows fit the one-class models. "
        "The held-out split is not used to fit parameters. Precision measures the anomaly correctness of predictions; "
        "recall measures recovered true anomalies; F1 is their harmonic mean; false-positive rate measures normal rows "
        "incorrectly flagged; detection time is measured in milliseconds for the held-out batch."
    )
    heading(document, "5.3 Algorithms and explainability", 2)
    add_bullets(document, [
        "Isolation Forest isolates rare points through random partitions.",
        "One-Class SVM learns a boundary around normal training behavior.",
        "Local Outlier Factor compares local density and is used in novelty mode.",
        "PyOD combination.average combines normalized model scores; bounds and threshold are learned during training.",
        "Rules retain direct evidence thresholds. ML detections state model version and normalized anomaly score.",
    ])

    heading(document, "6. Implementation")
    add_table(document, ["Layer", "Responsibilities"], [
        ("Routes/API", "Page authorization, JSON v1 endpoints, JWT cookies, SSE stream"),
        ("Services", "Upload lifecycle, shared processing, generator thread, trust score, PDF reports"),
        ("Engines", "Parsing, normalization, features, rules, PyOD ensemble, MITRE, correlation"),
        ("Repositories", "Raw parameterized SQL and transaction boundaries"),
        ("Knowledge base", "16-entry taxonomy and MITRE mapping"),
        ("Frontend", "API-driven values, EventSource updates, chart highlighting, static labels"),
    ])
    heading(document, "6.1 Analyst feedback", 2)
    document.add_paragraph(
        "Risk labels are confirmed or false_positive with reviewer and timestamp. Retraining re-extracts the shared "
        "features, saves a versioned joblib artifact, compares old/new F1, and records the run. Generator truth simulates "
        "analyst adjudication and is disclosed as a limitation."
    )

    heading(document, "7. Evaluation and Results")
    add_table(document, ["Model", "Precision", "Recall", "F1", "FPR", "Detection ms"], [
        (row["model"], f"{float(row['precision']):.3f}", f"{float(row['recall']):.3f}",
         f"{float(row['f1']):.3f}", f"{float(row['false_positive_rate']):.3f}",
         f"{float(row['detection_time_ms']):.2f}") for row in metric_rows
    ])
    document.add_paragraph(
        "Rules perform perfectly on this deliberately separable four-mode generator, while ML demonstrates a "
        "general anomaly interface and measurable false positives. These controlled results must not be presented as "
        "production accuracy. The ensemble provides cross-model behavior rather than a guaranteed improvement over every component."
    )
    add_figure(document, "confusion_matrix.png", "Figure 6. Combined ensemble confusion matrix", 5.3)
    add_figure(document, "roc_curve.png", "Figure 7. Combined ensemble ROC curve", 5.3)

    heading(document, "8. Security, Trust, and Limitations")
    heading(document, "8.1 Security controls", 2)
    add_bullets(document, [
        "No secret is hardcoded or returned by an API; .env is excluded from Git and Docker context.",
        "psycopg2 parameter binding separates SQL text from untrusted values.",
        "Flask-Login ownership checks protect page/data access; JWTs use HttpOnly SameSite cookies and refresh CSRF.",
        "Auth and generator controls are rate limited; CORS accepts only the configured origin.",
        "SHA-256 distinguishes modified evidence; source_status differentiates upload and generated rows.",
    ])
    heading(document, "8.2 Forensic Trust Score", 2)
    document.add_paragraph(
        "Trust = 0.40 × data quality + 0.30 × average detection risk + 0.30 × rule/ML binary agreement. "
        "The score describes evidence actionability, not system safety, and exposes all components in the API."
    )
    heading(document, "8.3 Limitations", 2)
    add_bullets(document, [
        "Synthetic modes are controlled and do not reproduce all real-world drift, ambiguity, or adversarial adaptation.",
        "Feedback labels are simulated stand-ins; additional labels may improve, reduce, or leave F1 unchanged.",
        "A process-local SSE broker and generator intentionally require one Gunicorn process; horizontal scale needs an external queue.",
        "MITRE mappings are hypotheses requiring analyst confirmation, not proof of attacker intent.",
        "The model artifact is runtime state and must be generated/retrained after a fresh deployment.",
    ])

    heading(document, "9. Testing and Deployment")
    document.add_paragraph(
        "The final suite uses a uniquely named temporary PostgreSQL database, applies all migrations, and removes only "
        "that database after testing. Coverage includes four formats, cleanup, hashing, taxonomy integrity, PyOD combination, "
        "feedback cycles, JWT/CORS, background generation, SSE, trust formula, design tokens, and Docker packaging."
    )
    add_table(document, ["Gate", "Fresh evidence"], [
        ("Automated suite", "Final unittest discovery count recorded in Phase 13 validation; zero failures"),
        ("Migration", "Fresh, idempotent, and populated-data validation"),
        ("Docker", "Clean Compose build; PostgreSQL healthy; Gunicorn serving"),
        ("Container workflow", "Login/dashboard 200; generator risk rows 0→10; stop; zero upload files"),
    ])
    heading(document, "9.1 Reproduction commands", 2)
    for command in (
        "python -m unittest discover -s tests -v", "python scripts/evaluate_models.py",
        "python scripts/generate_report_artifacts.py", "python scripts/build_final_report.py",
        "python scripts/run_migrations.py", "docker compose up --build",
    ):
        paragraph = document.add_paragraph()
        run = paragraph.add_run(command)
        run.font.name = "Consolas"
        run.font.size = Pt(9)

    heading(document, "10. Conclusion and Future Work")
    document.add_paragraph(
        "The final system satisfies the corrected architecture: schema-matched synthetic evaluation, explainable rules "
        "and ML, raw PostgreSQL repositories, controlled evidence storage, live in-memory generation, honest feedback "
        "measurement, API-driven real values, SSE updates, integrity features, and reproducible Docker deployment. "
        "Future work should evaluate anonymized real application logs, introduce drift monitoring and external event queues, "
        "and conduct analyst usability studies without changing the evidence/provenance contract."
    )

    heading(document, "References")
    references = [
        "MITRE ATT&CK. Enterprise techniques. https://attack.mitre.org/",
        "OWASP Foundation. Web Application Security guidance. https://owasp.org/",
        "PyOD documentation and source. https://pyod.readthedocs.io/",
        "scikit-learn User Guide. https://scikit-learn.org/stable/user_guide.html",
        "Flask documentation. https://flask.palletsprojects.com/",
        "PostgreSQL documentation. https://www.postgresql.org/docs/",
        "NIST. Secure Hash Standard (SHA-256), FIPS PUB 180-4.",
    ]
    add_numbered(document, references)

    heading(document, "Appendix A — Final Phase Evidence")
    phase_deliverables = [
        "Generator-shaped ML evaluation", "Raw PostgreSQL migrations/repositories", "Taxonomy/classification",
        "Four-format upload pipeline", "Live generator", "Analyst feedback", "JWT/REST hardening",
        "API-driven real values", "SSE/live UI", "Hash/trust", "Design tokens", "Docker deployment",
        "Diagrams/evaluation/Word report",
    ]
    add_table(document, ["Phase", "Principal deliverable", "Evidence record"], [
        (phase, phase_deliverables[phase - 1], f"docs/phase_{phase:02d}_validation.txt")
        for phase in range(1, 14)
    ])
    heading(document, "Appendix B — Principal API Endpoints")
    add_table(document, ["Method", "Endpoint", "Purpose"], [
        ("POST", "/api/v1/auth/login", "Cookie JWT login"), ("GET", "/api/v1/uploads", "Evidence history"),
        ("GET", "/api/v1/uploads/{id}", "Detailed evidence/risks/trust"),
        ("POST", "/api/v1/generator/start", "Start live mode"), ("POST", "/api/v1/generator/stop", "Stop live mode"),
        ("GET", "/api/v1/stream/live", "Committed risk-event SSE"),
        ("POST", "/api/v1/risk-events/{id}/label", "Analyst feedback"),
        ("GET", "/api/v1/model/labels-vs-f1", "Retraining behavior series"),
    ])
    heading(document, "Appendix C — Submission Review")
    add_bullets(document, [
        "Replace all cover-page placeholders and obtain required signatures.",
        "Update institution-specific page numbering, certificate wording, and formatting rules.",
        "Open the document in Word and update any institution-required table of contents.",
        "Re-run the artifact/report scripts after any final evaluation change.",
        "Do not claim synthetic metrics as production performance.",
    ])

    document.core_properties.title = "Intelligent Log Forensics — Final Project Report"
    document.core_properties.subject = "Explainable log anomaly detection and forensic integrity"
    document.core_properties.author = "Project report generator (student details pending)"
    document.core_properties.keywords = "log forensics, PyOD, Flask, PostgreSQL, MITRE ATT&CK"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(f"Generated {OUTPUT} ({OUTPUT.stat().st_size} bytes)")


if __name__ == "__main__":
    build()
